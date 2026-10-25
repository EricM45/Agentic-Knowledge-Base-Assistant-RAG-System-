"""Load Markdown, text, PDF, and Word files into sections."""

from __future__ import annotations

from pathlib import Path

from ragbot.nlp import clean_prose, markdown_sections
from ragbot.types import Document

SUPPORTED = {".txt", ".md", ".markdown", ".pdf", ".docx"}


def collect_files(paths: list[Path]) -> list[Path]:
    found: list[Path] = []
    for path in paths:
        if path.is_dir():
            found.extend(child for child in path.rglob("*") if child.is_file())
        elif path.is_file():
            found.append(path)
    unique: list[Path] = []
    seen: set[str] = set()
    for path in sorted(found, key=lambda item: str(item).lower()):
        if path.suffix.lower() not in SUPPORTED or _hidden(path):
            continue
        resolved = str(path.resolve())
        if resolved in seen:
            continue
        seen.add(resolved)
        unique.append(path)
    return unique


def load_paths(paths: list[Path]) -> tuple[list[Document], tuple[str, ...]]:
    documents: list[Document] = []
    skipped: list[str] = []
    for path in collect_files(paths):
        try:
            loaded = load_file(path)
        except Exception as exc:
            skipped.append(f"{path.name}: could not read ({exc.__class__.__name__})")
            continue
        if not loaded:
            reason = f"{path.name}: no extractable text"
            if path.suffix.lower() == ".pdf":
                reason += " (scanned pages need OCR)"
            skipped.append(reason)
            continue
        documents.extend(loaded)
    return documents, tuple(skipped)


def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown"}:
        return _from_sections(path, markdown_sections(path.read_text(encoding="utf-8-sig", errors="replace")))
    if suffix == ".txt":
        text = clean_prose(path.read_text(encoding="utf-8-sig", errors="replace"))
        if not text:
            return []
        return [Document(text=text, source=path.name, section="", page=None)]
    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix == ".docx":
        return _load_docx(path)
    return []


def _from_sections(path: Path, sections: list[tuple[str, str]]) -> list[Document]:
    documents: list[Document] = []
    for section, body in sections:
        text = body.strip()
        if text:
            documents.append(Document(text=text, source=path.name, section=section, page=None))
    return documents


def _load_pdf(path: Path) -> list[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    documents: list[Document] = []
    for number, page in enumerate(reader.pages, start=1):
        text = clean_prose(page.extract_text() or "")
        if text:
            documents.append(Document(text=text, source=path.name, section=f"Page {number}", page=number))
    return documents


def _load_docx(path: Path) -> list[Document]:
    from docx import Document as DocxDocument
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = DocxDocument(str(path))
    sections: list[tuple[str, list[str]]] = [("", [])]
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style
            style_name = style.name if style is not None and style.name else ""
            if style_name.startswith("Heading"):
                sections.append((text, []))
            else:
                sections[-1][1].append(text)
        elif isinstance(child, CT_Tbl):
            rows = _table_rows(Table(child, document))
            if rows:
                sections[-1][1].append("\n".join(rows))
    loaded: list[Document] = []
    for section, lines in sections:
        text = clean_prose("\n".join(lines))
        if text:
            loaded.append(Document(text=text, source=path.name, section=section, page=None))
    return loaded


def _table_rows(table) -> list[str]:
    rows: list[str] = []
    for row in table.rows:
        cells: list[str] = []
        seen: set[str] = set()
        for cell in row.cells:
            value = " ".join(cell.text.split())
            if value and value not in seen:
                seen.add(value)
                cells.append(value)
        if cells:
            rows.append(" | ".join(cells))
    return rows


def _hidden(path: Path) -> bool:
    return any(part.startswith(".") for part in path.parts)
