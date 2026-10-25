"""Unit tests for retrieval, grounding, and the sample knowledge base.

These tests use a hashed bag-of-words embedder so they do not download models.
"""

from __future__ import annotations

import hashlib
import sys
import tempfile
import unittest
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from ragbot.bm25 import BM25Index
from ragbot.chunker import chunk_documents, chunk_text
from ragbot.config import Settings
from ragbot.evaluate import problems
from ragbot.gold import FOLLOWUPS, GOLD
from ragbot.loader import load_file, load_paths
from ragbot.nlp import analyze_query, standalone_query, stem, tokenize
from ragbot.pipeline import RagPipeline
from ragbot.pinecone_store import sync_matches
from ragbot.retriever import CorpusIndex, maximal_marginal_relevance, reciprocal_rank_fusion
from ragbot.store import corpus_signature
from ragbot.types import Chunk, Hit

SAMPLE = ROOT / "sample_documents"


class HashEmbedder:
    model_name = "hash-bow-v1"
    dimension = 384

    def embed_documents(self, texts: list[str]) -> np.ndarray:
        if not texts:
            return np.zeros((0, self.dimension), dtype=np.float32)
        vectors = np.zeros((len(texts), self.dimension), dtype=np.float32)
        for row, text in enumerate(texts):
            for token in tokenize(text):
                bucket = int(hashlib.md5(token.encode("utf-8")).hexdigest()[:8], 16) % self.dimension
                vectors[row, bucket] += 1.0
            norm = float(np.linalg.norm(vectors[row]))
            if norm:
                vectors[row] /= norm
        return vectors

    def embed_query(self, text: str) -> np.ndarray:
        return self.embed_documents([text])[0]


def make_pipeline(directory: Path) -> RagPipeline:
    settings = Settings(
        embedding_model=HashEmbedder.model_name,
        index_dir=directory / "index",
        sample_dir=SAMPLE,
        backend="local-nlp",
        use_reranker=False,
        min_dense_score=1.1,
    )
    return RagPipeline(settings, embedder=HashEmbedder())


class NlpTests(unittest.TestCase):
    def test_stem_and_numbers(self) -> None:
        self.assertEqual(stem("extended"), "extend")
        self.assertEqual(stem("extend"), "extend")
        self.assertEqual(stem("days"), "day")
        self.assertEqual(stem("receives"), "receive")
        self.assertEqual(stem("costs"), "cost")
        self.assertIn("13.5", tokenize("stores 13.5 kWh"))
        self.assertIn("kwh", tokenize("stores 13.5 kWh"))
        self.assertNotIn("the", tokenize("The standard warranty is 12 years."))

    def test_follow_up_rewrite_keeps_standalone_questions(self) -> None:
        history = [{"role": "user", "content": "What is the standard warranty on the Helios Arc inverter?"}]
        rewritten = standalone_query("How long can it be extended?", history)
        self.assertIn("Arc", rewritten)
        self.assertIn("extended", rewritten)
        terms, extras = analyze_query("How many PTO days does a new employee receive?")
        self.assertIn("pto", terms)
        self.assertIn("paid", extras)
        self.assertIn("time", extras)
        untouched = standalone_query("What is the monthly price of Gridline monitoring?", history)
        self.assertEqual(untouched, "What is the monthly price of Gridline monitoring?")


class RetrievalTests(unittest.TestCase):
    def test_bm25_prefers_the_rare_term(self) -> None:
        index = BM25Index(
            [
                "The warehouse door uses alpha codes.",
                "The standard warranty on the Helios Arc hybrid inverter is 12 years.",
                "Community workshops happen for neighbors.",
            ]
        )
        scores = index.score("What is the Arc warranty?")
        self.assertEqual(int(np.argmax(scores)), 1)

    def test_reciprocal_rank_fusion_rewards_agreement(self) -> None:
        fused = [doc_id for doc_id, _score in reciprocal_rank_fusion([[0, 1, 2], [2, 0, 1]])]
        self.assertEqual(fused, [0, 2, 1])

    def test_mmr_skips_a_near_duplicate(self) -> None:
        vectors = np.array([[1.0, 0.0], [0.99, 0.14], [0.0, 1.0]], dtype=np.float32)
        vectors /= np.linalg.norm(vectors, axis=1, keepdims=True)
        relevance = np.array([0.95, 0.94, 0.55], dtype=np.float32)
        chosen = maximal_marginal_relevance(vectors, relevance, top_k=2, lambda_mult=0.2)
        self.assertIn(0, chosen)
        self.assertIn(2, chosen)

    def test_external_dense_hits_supply_the_similarity_scores(self) -> None:
        chunks = chunk_documents(load_paths([SAMPLE / "products.md"])[0], 800, 140)
        embeddings = HashEmbedder().embed_documents([chunk.text for chunk in chunks])
        index = CorpusIndex(chunks, embeddings)
        query = HashEmbedder().embed_query("Arc warranty")
        hits = index.search("Arc warranty", query, limit=len(chunks), dense_hits=[(0, 0.2), (1, 0.91)])
        scores = {hit.chunk.index: hit.dense_score for hit in hits}
        self.assertAlmostEqual(scores[0], 0.2, places=5)
        self.assertAlmostEqual(scores[1], 0.91, places=5)
        self.assertTrue(sync_matches({"signature": "abc", "index_name": "rag", "namespace": "abc"}, "abc", "rag"))
        self.assertFalse(sync_matches({"signature": "abc", "index_name": "rag", "namespace": "abc"}, "xyz", "rag"))

    def test_short_text_is_one_chunk(self) -> None:
        self.assertEqual(chunk_text("Only one sentence lives here.", 200, 40), ["Only one sentence lives here."])

    def test_long_text_overlaps_on_a_sentence(self) -> None:
        sentences = [
            f"Sentence {index:02d} carries marker M{index:02d} and has enough words to occupy space in the chunk."
            for index in range(10)
        ]
        chunks = chunk_text(" ".join(sentences), size=180, overlap=100)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(chunk) <= 180 or len(chunk.split()) == 1 for chunk in chunks))
        first_sentences = [part for part in chunks[0].split(". ") if part]
        self.assertTrue(any(part in chunks[1] for part in first_sentences))


class CorpusTests(unittest.TestCase):
    def test_sample_sections_keep_the_warranty_sentence(self) -> None:
        documents, skipped = load_paths([SAMPLE])
        self.assertEqual(skipped, ())
        chunks = chunk_documents(documents, 800, 140)
        self.assertGreater(len(chunks), 8)
        self.assertTrue(any("12 years" in chunk.text and "products.md" == chunk.source for chunk in chunks))

    def test_docx_headings_become_sections(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "leave.docx"
            document = Document()
            document.add_heading("Leave", level=1)
            document.add_paragraph("Primary caregivers receive 16 weeks of paid parental leave.")
            document.save(str(path))
            loaded = load_file(path)
        self.assertTrue(any("16 weeks" in item.text and item.section == "Leave" for item in loaded))

    def test_signature_changes_when_a_file_changes(self) -> None:
        settings = Settings(embedding_model="hash-bow-v1")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "note.md"
            path.write_text("# Note\n\nThe alpha code is 4242 for the warehouse door.\n", encoding="utf-8")
            first = corpus_signature([path], settings, "hash-bow-v1")
            path.write_text("# Note\n\nThe alpha code is 9999 for the warehouse door.\n", encoding="utf-8")
            second = corpus_signature([path], settings, "hash-bow-v1")
        self.assertNotEqual(first, second)


class PipelineTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.pipeline = make_pipeline(Path(self._tmp.name))
        self.pipeline.index_paths([SAMPLE])

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_gold_questions(self) -> None:
        for item in GOLD:
            answer = self.pipeline.ask(item["question"])
            self.assertEqual(problems(answer, item), [], msg=f"{item['question']}\n{answer.text}")

    def test_follow_up_uses_the_earlier_topic(self) -> None:
        for item in FOLLOWUPS:
            answer = self.pipeline.ask(item["question"], history=item["history"])
            self.assertEqual(problems(answer, item), [], msg=answer.text)
            self.assertIn("[1]", answer.text)

    def test_saved_index_reloads(self) -> None:
        restored = RagPipeline(self.pipeline.settings, embedder=HashEmbedder())
        self.assertTrue(restored.load_cached([SAMPLE]))
        self.assertEqual(restored.stats.chunks, self.pipeline.stats.chunks)
        answer = restored.ask("Who is the chief executive of Helios Labs?")
        self.assertIn("Amina Shah", answer.text)

    def test_custom_generator_is_used_when_the_passage_is_grounded(self) -> None:
        class Stub:
            def generate(self, question, hits, history, settings, idf, backend):
                from ragbot.types import Generated

                return Generated(text="The standard warranty is 12 years [1].", backend="stub", refused=False)

        self.pipeline.generator = Stub()
        answer = self.pipeline.ask("What is the standard warranty on the Helios Arc inverter?")
        self.assertEqual(answer.backend, "stub")
        self.assertIn("12", answer.text)
        self.assertTrue(answer.sources)


class PromptTests(unittest.TestCase):
    def test_prompt_includes_the_passage_and_the_question(self) -> None:
        from ragbot.generator import build_context_prompt

        chunk = Chunk(
            index=0,
            text="The standard warranty on the Helios Arc hybrid inverter is 12 years.",
            source="products.md",
            section="Arc inverter",
            page=None,
        )
        hit = Hit(chunk=chunk, dense_score=0.5, sparse_score=2.0, fused_score=0.03)
        prompt = build_context_prompt("What is the warranty?", [hit], [])
        self.assertIn("[1]", prompt)
        self.assertIn("12 years", prompt)
        self.assertIn("What is the warranty?", prompt)
        self.assertNotIn("Paris", prompt)


class MemoryVectorStore:
    def __init__(self) -> None:
        self.sync_count = 0
        self.namespaces: dict[str, np.ndarray] = {}

    def sync(self, chunks, embeddings, namespace: str, previous_namespace: str | None = None) -> None:
        del chunks, previous_namespace
        self.sync_count += 1
        self.namespaces[namespace] = np.asarray(embeddings, dtype=np.float32)

    def query(self, vector, namespace: str, top_k: int) -> list[tuple[int, float]]:
        matrix = self.namespaces[namespace]
        scores = matrix @ np.asarray(vector, dtype=np.float32)
        order = np.argsort(-scores)[:top_k]
        return [(int(row), float(scores[row])) for row in order]


class PineconePathTests(unittest.TestCase):
    def test_remote_vectors_are_uploaded_once_and_used_for_search(self) -> None:
        store = MemoryVectorStore()
        with tempfile.TemporaryDirectory() as tmp:
            settings = Settings(
                embedding_model=HashEmbedder.model_name,
                index_dir=Path(tmp) / "index",
                sample_dir=SAMPLE,
                backend="local-nlp",
                use_reranker=False,
                min_dense_score=1.1,
                vector_store="pinecone",
                pinecone_index="rag-chatbot",
            )
            pipeline = RagPipeline(settings, embedder=HashEmbedder(), vector_store=store)
            pipeline.index_paths([SAMPLE])
            self.assertEqual(store.sync_count, 1)
            self.assertEqual(pipeline.vector_store_label(), "Pinecone")
            answer = pipeline.ask("What is the standard warranty on the Helios Arc inverter?")
            self.assertEqual(store.sync_count, 1)
            self.assertEqual(answer.vector_store, "pinecone")
            self.assertIn("12", answer.text)
            self.assertFalse(answer.refused)


if __name__ == "__main__":
    unittest.main()
